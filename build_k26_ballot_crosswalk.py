#!/usr/bin/env python3
"""
Pair K25 and K26 ballots for the 25→26 Sankey, and preview K25 pin corrections.

By default writes only:

  data/k25_k26_ballot_union.csv
  data/k25_coord_corrections_preview.csv

Flags that rewrite site/data/station_coordinates.json:

  --apply-geo          move clearly wrong pins and write official K26 ITM
  --sync-keys          copy the better pin onto name|N and name|N.0 twins
  --backfill-missing   copy the paired K26 point onto K25 keys with no pin

Pairing policy (a K25 ballot is used at most once):

  id / id_dot1
    Settlement code + ballot number, with the usual .1 → base fallback.
    Kept when venue names agree, when either name has no distinctive tokens,
    or when both names disagree but the pins are within 150 m (same building,
    renamed). Rejected when the names disagree and the pins are farther than
    150 m or one side has no pin — those numbers were reused.

  venue
    Among leftovers, one K26 cluster whose distinctive venue-name tokens equal
    exactly one K25 venue in that settlement. Ballots are zipped in ballot-number
    order up to the shorter side. Two K26 clusters with the same name are skipped.

  coord
    Among leftovers, a K26 cluster and a K25 venue-grade point (not a town
    center) that are mutual nearest neighbors within 150 m, and whose names do
    not conflict. Same zip as venue.

Coordinate preview (not applied):
  Accepted pairs whose venue names agree. google_venue / venue pins move only
  when they are at least 400 m off. Town-center and missing pins are listed
  only when the official point is at least 400 m away. Manual pins are never
  proposed.
"""

import csv
import json
import math
import re
from collections import defaultdict
from pathlib import Path

import openpyxl
import pandas as pd
from pyproj import Transformer

from party_config import ELECTIONS

ROOT = Path(__file__).resolve().parent
STATIONS_XLSX = ROOT / 'data' / 'k26_cec' / 'stations.xlsx'
CLUSTERS_XLSX = ROOT / 'data' / 'k26_cec' / 'clusters.xlsx'
COORDS_PATH = ROOT / 'site' / 'data' / 'station_coordinates.json'
LOCATIONS_25 = ROOT / 'data' / 'ballot_locations_25.json'
UNION_PATH = ROOT / 'data' / 'k25_k26_ballot_union.csv'
PREVIEW_PATH = ROOT / 'data' / 'k25_coord_corrections_preview.csv'

GENERIC = {
    'בית', 'ספר', 'ביס', 'ביהס', 'גן', 'ילדים', 'עירוני', 'ממלכתי', 'דתי',
    'תיכון', 'חטיבה', 'חטיבת', 'ביניים', 'עליונה', 'יסודי', 'ממד', 'מועדון',
    'מרכז', 'קהילתי', 'מתנס', 'אולם', 'כנסת', 'מקיף', 'על', 'שם', 'עש',
    'לשעבר', 'בנים', 'בנות',
}
CORRECTABLE = {'google_venue', 'venue', 'settlement', 'not_found'}
VENUE_GRADE = {'google_venue', 'venue', 'manual', 'manual_correction', 'manual_override', 'google_venue_haifa'}
MOVE_MIN_M = 400

TO_WGS = Transformer.from_crs('EPSG:2039', 'EPSG:4326', always_xy=True)


def norm_ballot(b):
    if b is None:
        return None
    b = str(b).strip()
    if b.endswith('.0'):
        b = b[:-2]
    return b


def ballot_sort_key(b):
    text = norm_ballot(b)
    try:
        return (0, float(text))
    except (TypeError, ValueError):
        return (1, text or '')


def fold(s):
    s = str(s or '')
    s = s.replace('״', '"').replace('׳', "'")
    s = s.replace('בי"ס', ' ').replace("בי'ס", ' ').replace('ביה"ס', ' ').replace('ביהס', ' ')
    s = re.sub(r'["\'\-–.,()]', ' ', s)
    s = re.sub(r'\s+', ' ', s).strip()
    return s


def tokens(s):
    return [t for t in fold(s).split() if len(t) > 1 and t not in GENERIC]


def signature(s):
    return frozenset(tokens(s))


def hav(a, b):
    radius = 6371000.0
    p1, p2 = math.radians(a[0]), math.radians(b[0])
    dp, dl = math.radians(b[0] - a[0]), math.radians(b[1] - a[1])
    x = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * radius * math.asin(math.sqrt(min(1.0, x)))


def load_k25():
    df = pd.read_csv(ROOT / 'ballot25.csv', encoding='utf-8-sig')
    df = df[df['סמל ישוב'] != 9999].copy()
    locs = json.loads(LOCATIONS_25.read_text(encoding='utf-8'))['settlements']
    loc_by = {}
    for code, settlement in locs.items():
        if not isinstance(settlement, dict):
            continue
        for ballot in settlement.get('ballots', []):
            loc_by[(int(code), norm_ballot(ballot.get('ballot')))] = ballot.get('location') or ''

    coords = json.loads(COORDS_PATH.read_text(encoding='utf-8'))['stations']
    by_name = {}
    for key, val in coords.items():
        if '|' not in key:
            continue
        name, bal = key.rsplit('|', 1)
        by_name.setdefault((name, norm_ballot(bal)), []).append((key, val))

    rows = []
    for _, rec in df.iterrows():
        code = int(rec['סמל ישוב'])
        ballot = norm_ballot(rec['קלפי'])
        name = str(rec['שם ישוב']).strip()
        venue = loc_by.get((code, ballot), '')
        hits = by_name.get((name, ballot), [])
        lat = lng = source = coord_key = stored_location = None
        for key, val in hits:
            if val.get('lat') is not None and lat is None:
                lat, lng = float(val['lat']), float(val['lng'])
                source = val.get('source')
                coord_key = key
                stored_location = val.get('location') or ''
            elif lat is None and coord_key is None:
                source = val.get('source')
                coord_key = key
                stored_location = val.get('location') or ''
        rows.append({
            'code': code,
            'name': name,
            'ballot': ballot,
            'bid': f'{code}__{ballot}',
            'eligible': int(rec['בזב'] or 0),
            'valid': int(rec['כשרים'] or 0),
            'venue': venue,
            'sig': signature(venue),
            'lat': lat,
            'lng': lng,
            'source': source,
            'coord_key': coord_key,
            'stored_location': stored_location or '',
        })
    return rows


def load_k26():
    wb = openpyxl.load_workbook(CLUSTERS_XLSX, read_only=True, data_only=True)
    crows = list(wb.active.iter_rows(values_only=True))
    wb.close()
    ch = {n: i for i, n in enumerate(crows[0])}
    cluster_xy = {}
    for rec in crows[1:]:
        if rec[ch['סמל ישוב בחירות']] is None or rec[ch['X']] is None or rec[ch['Y']] is None:
            continue
        code = int(rec[ch['סמל ישוב בחירות']])
        cid = rec[ch['סמל רכוז']]
        lng, lat = TO_WGS.transform(float(rec[ch['X']]), float(rec[ch['Y']]))
        cluster_xy[(code, cid)] = (lat, lng)

    wb = openpyxl.load_workbook(STATIONS_XLSX, read_only=True, data_only=True)
    srows = list(wb.active.iter_rows(values_only=True))
    wb.close()
    sh = {n: i for i, n in enumerate(srows[0])}
    rows = []
    for rec in srows[1:]:
        if rec[sh['סמל ישוב בחירות']] is None or rec[sh['סמל קלפי']] is None:
            continue
        code = int(rec[sh['סמל ישוב בחירות']])
        ballot = norm_ballot(rec[sh['סמל קלפי']])
        cid = rec[sh['סמל רכוז']]
        xy = cluster_xy.get((code, cid))
        venue = str(rec[sh['מקום קלפי']] or '').strip()
        rows.append({
            'code': code,
            'name': str(rec[sh['שם ישוב בחירות']] or '').strip(),
            'ballot': ballot,
            'bid': f'{code}__{ballot}',
            'eligible': int(rec[sh['בוחרי כנסת']] or 0),
            'venue': venue,
            'address': str(rec[sh['כתובת קלפי']] or '').strip(),
            'cluster': cid,
            'sig': signature(venue),
            'lat': xy[0] if xy else None,
            'lng': xy[1] if xy else None,
        })
    return rows


def pair_distance(a, b):
    if a['lat'] is None or b['lat'] is None:
        return None
    return hav((a['lat'], a['lng']), (b['lat'], b['lng']))


def id_conflict(k25, k26):
    """True when the shared ballot number is not the same place."""
    if not k25['sig'] or not k26['sig']:
        return False
    if k25['sig'] == k26['sig']:
        return False
    dist = pair_distance(k25, k26)
    if dist is not None and dist <= 150:
        return False
    return True


def zip_ballots(k25_list, k26_list, method):
    a = sorted(k25_list, key=lambda r: ballot_sort_key(r['ballot']))
    b = sorted(k26_list, key=lambda r: ballot_sort_key(r['ballot']))
    return [(method, left, right) for left, right in zip(a, b)]


def build_pairs(k25, k26):
    k25_by = {r['bid']: r for r in k25}
    k26_ids = {r['bid'] for r in k26}
    accepted = []
    rejected = []
    used25, used26 = set(), set()

    for right in k26:
        bid = right['bid']
        kind, left_id = None, None
        if bid in k25_by:
            kind, left_id = 'id', bid
        elif bid.split('__')[1].endswith('.1'):
            base = bid.split('__')[0] + '__' + bid.split('__')[1].split('.')[0]
            if base not in k26_ids and base in k25_by:
                kind, left_id = 'id_dot1', base
        if kind is None:
            continue
        left = k25_by[left_id]
        if id_conflict(left, right):
            rejected.append((kind, left, right))
            continue
        accepted.append((kind, left, right))
        used25.add(left['bid'])
        used26.add(right['bid'])

    # Venue-name groups among leftovers.
    k25_groups = defaultdict(list)
    for left in k25:
        if left['bid'] in used25 or not left['sig']:
            continue
        k25_groups[(left['code'], left['sig'])].append(left)
    k26_groups = defaultdict(list)
    for right in k26:
        if right['bid'] in used26 or not right['sig']:
            continue
        k26_groups[(right['code'], right['sig'])].append(right)

    sigs_25 = defaultdict(list)
    sigs_26 = defaultdict(list)
    for (code, sig), ballots in k25_groups.items():
        sigs_25[(code, sig)].append(ballots)
    for (code, sig), ballots in k26_groups.items():
        clusters = {b['cluster'] for b in ballots}
        sigs_26[(code, sig)].append((clusters, ballots))

    for key, groups26 in sigs_26.items():
        groups25 = sigs_25.get(key)
        if not groups25 or len(groups25) != 1 or len(groups26) != 1:
            continue
        clusters, ballots26 = groups26[0]
        if len(clusters) != 1:
            continue
        for kind, left, right in zip_ballots(groups25[0], ballots26, 'venue'):
            accepted.append((kind, left, right))
            used25.add(left['bid'])
            used26.add(right['bid'])

    # Mutual nearest venue-grade point, within 150 m, names not in conflict.
    points = defaultdict(list)  # (code, rounded lat, lng) -> k25 ballots
    for left in k25:
        if left['bid'] in used25 or left['lat'] is None or left['source'] not in VENUE_GRADE:
            continue
        points[(left['code'], round(left['lat'], 4), round(left['lng'], 4))].append(left)

    clusters = defaultdict(list)
    for right in k26:
        if right['bid'] in used26 or right['lat'] is None:
            continue
        clusters[(right['code'], right['cluster'])].append(right)

    point_items = []
    for (code, lat_r, lng_r), ballots in points.items():
        lat = sum(b['lat'] for b in ballots) / len(ballots)
        lng = sum(b['lng'] for b in ballots) / len(ballots)
        point_items.append((code, (lat, lng), ballots))

    cluster_items = []
    for (code, cid), ballots in clusters.items():
        cluster_items.append((code, cid, (ballots[0]['lat'], ballots[0]['lng']), ballots))

    by_code_pts = defaultdict(list)
    by_code_cls = defaultdict(list)
    for item in point_items:
        by_code_pts[item[0]].append(item)
    for item in cluster_items:
        by_code_cls[item[0]].append(item)

    def nearest_point(code, xy):
        best, best_d = None, None
        for item in by_code_pts.get(code, []):
            dist = hav(xy, item[1])
            if best_d is None or dist < best_d:
                best, best_d = item, dist
        return best, best_d

    def nearest_cluster(code, xy, skip_cid=None):
        best, best_d = None, None
        for item in by_code_cls.get(code, []):
            if item[1] == skip_cid:
                continue
            dist = hav(xy, item[2])
            if best_d is None or dist < best_d:
                best, best_d = item, dist
        return best, best_d

    for code, cid, xy, ballots26 in cluster_items:
        point, dist = nearest_point(code, xy)
        if point is None or dist > 150:
            continue
        back, back_d = nearest_cluster(code, point[1])
        if back is None or back[1] != cid:
            continue
        ballots25 = [b for b in point[2] if b['bid'] not in used25]
        ballots26 = [b for b in ballots26 if b['bid'] not in used26]
        if not ballots25 or not ballots26:
            continue
        sigs25 = {b['sig'] for b in ballots25 if b['sig']}
        sigs26 = {b['sig'] for b in ballots26 if b['sig']}
        if sigs25 and sigs26 and sigs25.isdisjoint(sigs26):
            continue
        for kind, left, right in zip_ballots(ballots25, ballots26, 'coord'):
            accepted.append((kind, left, right))
            used25.add(left['bid'])
            used26.add(right['bid'])

    return accepted, rejected, used25, used26


def correction_candidates(pairs):
    """Pins we would move. Not applied."""
    proposals = []
    seen = set()
    for method, left, right in pairs:
        if right['lat'] is None or not left['coord_key']:
            continue
        if left['source'] not in CORRECTABLE:
            continue
        if left['sig'] and right['sig'] and left['sig'] != right['sig']:
            continue
        if left['stored_location'] and signature(left['stored_location']) and right['sig']:
            if signature(left['stored_location']) != right['sig'] and left['sig'] != right['sig']:
                continue
        dist = pair_distance(left, right)
        if dist is None:
            dist = 10 ** 9
        if dist < MOVE_MIN_M:
            continue
        if left['coord_key'] in seen:
            continue
        seen.add(left['coord_key'])
        proposals.append({
            'coord_key': left['coord_key'],
            'settlement': left['name'],
            'ballot': left['ballot'],
            'old_source': left['source'] or '',
            'old_lat': left['lat'],
            'old_lng': left['lng'],
            'new_lat': round(right['lat'], 6),
            'new_lng': round(right['lng'], 6),
            'distance_m': None if dist >= 10 ** 8 else round(dist, 1),
            'match_method': method,
            'k25_venue': left['venue'],
            'k26_venue': right['venue'],
            'stored_location': left['stored_location'],
        })
    return proposals


def earlier_elections(proposals):
    """Which pre-K25 ballots share a coordinate key we would rewrite."""
    wanted = {}
    for row in proposals:
        wanted[(fold_name(row['settlement']), norm_ballot(row['ballot']))] = row['coord_key']
    hits = defaultdict(set)
    for election_id in ['16', '17', '18', '19', '20', '21', '22', '23', '24']:
        cfg = ELECTIONS[election_id]
        path = ROOT / cfg['file']
        if not path.exists():
            continue
        df = pd.read_csv(path, encoding=cfg['encoding'], usecols=lambda c: c in ('שם ישוב', cfg.get('ballot_field', 'קלפי'), 'סמל ישוב'))
        ballot_field = cfg.get('ballot_field', 'קלפי')
        divisor = cfg.get('ballot_number_divisor', 1)
        for name, ballot in zip(df['שם ישוב'], df[ballot_field]):
            text = norm_ballot(ballot)
            if divisor > 1:
                try:
                    number = int(text)
                    if number % divisor == 0:
                        text = str(number // divisor)
                except (TypeError, ValueError):
                    pass
            key = wanted.get((fold_name(str(name)), text))
            if key:
                hits[election_id].add(key)
    return hits


def fold_name(name):
    name = str(name or '')
    name = name.replace('-', ' ').replace('–', ' ')
    name = name.replace("'", '').replace('\u05f3', '').replace('"', '').replace('\u05f4', '')
    name = name.replace('(', ' ').replace(')', ' ')
    name = ' '.join(name.split()).replace('יי', 'י')
    return name


def write_union(k25, k26, pairs, rejected):
    rejected_ids = {right['bid'] for _, _, right in rejected}
    paired26 = {right['bid']: (i, method, left) for i, (method, left, right) in enumerate(pairs, 1)}
    paired25 = {left['bid'] for _, left, _ in pairs}
    fieldnames = [
        'pair_id', 'decision', 'rejected_id', 'distance_m',
        'k25_code', 'k25_name', 'k25_ballot', 'k25_eligible', 'k25_valid',
        'k25_venue', 'k25_lat', 'k25_lng', 'k25_coord_source',
        'k26_code', 'k26_name', 'k26_ballot', 'k26_eligible',
        'k26_venue', 'k26_address', 'k26_cluster', 'k26_lat', 'k26_lng',
    ]
    rows = []

    def blank25():
        return {k: '' for k in fieldnames if k.startswith('k25_')}

    def blank26():
        return {k: '' for k in fieldnames if k.startswith('k26_')}

    def put25(row, left):
        row.update({
            'k25_code': left['code'], 'k25_name': left['name'], 'k25_ballot': left['ballot'],
            'k25_eligible': left['eligible'], 'k25_valid': left['valid'], 'k25_venue': left['venue'],
            'k25_lat': '' if left['lat'] is None else round(left['lat'], 6),
            'k25_lng': '' if left['lng'] is None else round(left['lng'], 6),
            'k25_coord_source': left['source'] or '',
        })

    def put26(row, right):
        row.update({
            'k26_code': right['code'], 'k26_name': right['name'], 'k26_ballot': right['ballot'],
            'k26_eligible': right['eligible'], 'k26_venue': right['venue'],
            'k26_address': right['address'], 'k26_cluster': right['cluster'],
            'k26_lat': '' if right['lat'] is None else round(right['lat'], 6),
            'k26_lng': '' if right['lng'] is None else round(right['lng'], 6),
        })

    for i, (method, left, right) in enumerate(pairs, 1):
        dist = pair_distance(left, right)
        row = {'pair_id': f'p{i:05d}', 'decision': method, 'rejected_id': '',
               'distance_m': '' if dist is None else round(dist, 1)}
        put25(row, left)
        put26(row, right)
        rows.append(row)
    for left in k25:
        if left['bid'] in paired25:
            continue
        row = {'pair_id': '', 'decision': 'k25_only', 'rejected_id': '', 'distance_m': ''}
        put25(row, left)
        row.update(blank26())
        rows.append(row)
    for right in k26:
        if right['bid'] in paired26:
            continue
        row = {'pair_id': '', 'decision': 'k26_only', 'rejected_id': 'yes' if right['bid'] in rejected_ids else '',
               'distance_m': ''}
        row.update(blank25())
        put26(row, right)
        rows.append(row)

    with UNION_PATH.open('w', encoding='utf-8-sig', newline='') as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)
    return len(rows)


def write_preview(proposals, earlier):
    key_elections = defaultdict(list)
    for election_id, keys in earlier.items():
        for key in keys:
            key_elections[key].append(election_id)
    fieldnames = [
        'coord_key', 'settlement', 'ballot', 'old_source', 'old_lat', 'old_lng',
        'new_lat', 'new_lng', 'distance_m', 'match_method', 'k25_venue', 'k26_venue',
        'also_in_elections',
    ]
    with PREVIEW_PATH.open('w', encoding='utf-8-sig', newline='') as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames)
        writer.writeheader()
        for row in proposals:
            out = {k: row.get(k, '') for k in fieldnames}
            out['also_in_elections'] = ' '.join(key_elections.get(row['coord_key'], []))
            writer.writerow(out)


def band(dist):
    if dist is None:
        return 'no_pin'
    if dist < 150:
        return '<150m'
    if dist < 400:
        return '150-400m'
    if dist < 1000:
        return '400m-1km'
    if dist < 5000:
        return '1-5km'
    return '>5km'


def summarize(k25, k26, pairs, rejected, proposals, earlier):
    from collections import Counter
    k25_valid = sum(r['valid'] for r in k25)
    k25_elig = sum(r['eligible'] for r in k25)
    k26_elig = sum(r['eligible'] for r in k26)
    print('K25 ballots', len(k25), 'valid votes', f'{k25_valid:,}', 'eligible', f'{k25_elig:,}')
    print('K26 ballots', len(k26), 'eligible', f'{k26_elig:,}')
    print('pairs', len(pairs), 'rejected id', len(rejected))
    by = Counter(m for m, _, _ in pairs)
    for method in ('id', 'id_dot1', 'venue', 'coord'):
        subset = [(l, r) for m, l, r in pairs if m == method]
        if not subset:
            print(f'  {method:8} 0')
            continue
        v = sum(l['valid'] for l, _ in subset)
        e26 = sum(r['eligible'] for _, r in subset)
        dists = [pair_distance(l, r) for l, r in subset]
        known = [d for d in dists if d is not None]
        known.sort()
        med = round(known[len(known) // 2]) if known else None
        print(f'  {method:8} {len(subset):5}  k25 votes {v/k25_valid:6.1%}  k26 eligible {e26/k26_elig:6.1%}  median pin gap {med} m  n_with_both_pins {len(known)}')
        print('           gaps', dict(Counter(band(d) for d in dists)))
    v = sum(l['valid'] for _, l, _ in pairs)
    e26 = sum(r['eligible'] for _, _, r in pairs)
    b25 = len({l['bid'] for _, l, _ in pairs})
    b26 = len({r['bid'] for _, _, r in pairs})
    print(f'TOTAL pairs {len(pairs)}  k25 ballots {b25/len(k25):.1%}  k25 votes {v/k25_valid:.1%}  k26 ballots {b26/len(k26):.1%}  k26 eligible {e26/k26_elig:.1%}')
    rej_v = sum(l['valid'] for _, l, _ in rejected)
    rej_e = sum(r['eligible'] for _, _, r in rejected)
    print(f'Rejected id matches: {len(rejected)} ballots, k25 votes {rej_v:,} ({rej_v/k25_valid:.1%}), k26 eligible {rej_e:,} ({rej_e/k26_elig:.1%})')

    print('\nCorrection preview (not applied):', len(proposals))
    src = Counter(p['old_source'] for p in proposals)
    print('  by old source', dict(src))
    print('  by match', dict(Counter(p['match_method'] for p in proposals)))
    dists = [p['distance_m'] for p in proposals if p['distance_m'] is not None]
    dists.sort()
    if dists:
        print(f'  move distance median {dists[len(dists)//2]} m  p90 {dists[int(len(dists)*0.9)]} m  max {dists[-1]} m')
        print('  bands', dict(Counter(band(d) for d in dists)))
    any_earlier = 0
    for election_id in ['16', '17', '18', '19', '20', '21', '22', '23', '24']:
        n = len(earlier.get(election_id, ()))
        print(f'  keys also present in K{election_id}: {n}')
        any_earlier = max(any_earlier, n)
    keys_any = set()
    for keys in earlier.values():
        keys_any |= keys
    print(f'  keys shared with at least one of K16–K24: {len(keys_any)} of {len(proposals)}')
    far = sorted(proposals, key=lambda p: -(p['distance_m'] or 0))[:8]
    print('  farthest proposed moves:')
    for p in far:
        print(f"    {p['distance_m']} m  {p['settlement']} {p['ballot']}  {p['old_source']}  {p['k25_venue']}  →  {p['k26_venue']}")


def _point_ok(lat, lng):
    if lat is None or lng is None:
        return False
    if abs(lat - 26.06194) < 0.002 and abs(lng - 33.01166) < 0.02:
        return False
    return 29.4 < lat < 33.45 and 34.2 < lng < 35.95


def apply_geo(proposals, k25, k26):
    """Move clearly wrong pins, and store official K26 cluster coordinates.

    A move under 2 km is left alone. X=0,Y=0 clusters are not coordinates.
    A K26 point is not written onto a key whose K25 ballot is a different venue.
    """
    from generate_map_data import normalize_name

    clear_m = 2000
    data = json.loads(COORDS_PATH.read_text(encoding='utf-8'))
    stations = data['stations']
    moved = 0
    for row in proposals:
        dist = row['distance_m']
        if dist is None or dist < clear_m:
            continue
        if not _point_ok(row['new_lat'], row['new_lng']):
            continue
        entry = stations.get(row['coord_key'])
        if entry is None or entry.get('source') not in CORRECTABLE:
            continue
        entry['prev_source'] = entry.get('source')
        entry['prev_lat'] = entry.get('lat')
        entry['prev_lng'] = entry.get('lng')
        entry['lat'] = row['new_lat']
        entry['lng'] = row['new_lng']
        entry['source'] = 'cec_k26'
        moved += 1

    k25_sig = {(r['code'], r['ballot']): r['sig'] for r in k25}
    added = replaced = skipped = 0
    for station in k26:
        if not _point_ok(station['lat'], station['lng']):
            continue
        key = f"{normalize_name(station['name'])}|{station['ballot']}"
        other = k25_sig.get((station['code'], station['ballot']))
        if other and station['sig'] and other != station['sig']:
            skipped += 1
            continue
        entry = stations.get(key)
        if entry is None:
            stations[key] = {
                'settlement': normalize_name(station['name']),
                'ballot': station['ballot'],
                'location': station['venue'],
                'lat': round(station['lat'], 6),
                'lng': round(station['lng'], 6),
                'source': 'cec_k26',
            }
            added += 1
            continue
        if entry.get('source') in ('manual', 'manual_correction', 'manual_override'):
            continue
        old = (entry.get('lat'), entry.get('lng'))
        dist = hav(old, (station['lat'], station['lng'])) if _point_ok(*old) else clear_m
        if dist < 50 and entry.get('source') == 'cec_k26':
            continue
        entry['prev_source'] = entry.get('source')
        entry['prev_lat'] = entry.get('lat')
        entry['prev_lng'] = entry.get('lng')
        entry['lat'] = round(station['lat'], 6)
        entry['lng'] = round(station['lng'], 6)
        entry['source'] = 'cec_k26'
        if not entry.get('location'):
            entry['location'] = station['venue']
        replaced += 1

    from collections import Counter
    counts = Counter(v.get('source') for v in stations.values())
    data['stats'] = {
        'venue_match': counts.get('venue', 0) + counts.get('google_venue', 0) + counts.get('cec_k26', 0),
        'settlement_only': counts.get('settlement', 0),
        'not_found': counts.get('not_found', 0),
    }
    COORDS_PATH.write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(f'Coordinates: {moved} old pins moved ≥2 km, {added} K26 keys added, {replaced} K26 keys set to ITM, {skipped} conflicting numbers left alone')


def sync_ballot_key_twins():
    """Give both ``name|2`` and ``name|2.0`` the better pin.

    Older geocoding stored integer ballots as ``2.0``. Later writers stored
    ``2``. The K26 correction updated only the one key the crosswalk had
    recorded, so the twin could keep a wrong town-center pin.
    """
    priority = {
        'cec_k26': 6, 'manual_override': 5, 'manual_correction': 5, 'manual': 5,
        'google_venue_haifa': 4, 'google_venue': 3, 'venue': 2, 'settlement': 1,
        'not_found': 0,
    }
    data = json.loads(COORDS_PATH.read_text(encoding='utf-8'))
    stations = data['stations']
    groups = {}
    for key in stations:
        if '|' not in key:
            continue
        name, ballot = key.rsplit('|', 1)
        canon = ballot[:-2] if ballot.endswith('.0') else ballot
        groups.setdefault((name, canon), []).append(key)

    updated = groups_fixed = 0
    for keys in groups.values():
        if len(keys) < 2:
            continue
        ranked = sorted(keys, key=lambda k: (
            priority.get(stations[k].get('source'), 0),
            1 if _point_ok(stations[k].get('lat'), stations[k].get('lng')) else 0,
        ), reverse=True)
        winner = stations[ranked[0]]
        if not _point_ok(winner.get('lat'), winner.get('lng')):
            continue
        wlat, wlng = float(winner['lat']), float(winner['lng'])
        changed = False
        for key in ranked[1:]:
            entry = stations[key]
            olat, olng = entry.get('lat'), entry.get('lng')
            far = not _point_ok(olat, olng) or hav((float(olat), float(olng)), (wlat, wlng)) > 50
            if not far:
                continue
            entry['prev_source'] = entry.get('source')
            entry['prev_lat'] = olat
            entry['prev_lng'] = olng
            entry['lat'] = wlat
            entry['lng'] = wlng
            entry['source'] = winner.get('source')
            updated += 1
            changed = True
        if changed:
            groups_fixed += 1

    COORDS_PATH.write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(f'Synced {updated} duplicate ballot keys in {groups_fixed} ballots')


def backfill_missing_k25_pins():
    """Copy the paired K26 coordinate onto K25 keys that were never geocoded."""
    from generate_map_data import normalize_name

    data = json.loads(COORDS_PATH.read_text(encoding='utf-8'))
    stations = data['stations']

    def usable(entry):
        return entry is not None and _point_ok(entry.get('lat'), entry.get('lng'))

    def locate(name, ballot):
        name = normalize_name(name)
        for key in (f'{name}|{ballot}', f'{name}|{ballot}.0'):
            entry = stations.get(key)
            if usable(entry):
                return entry
        return None

    rows = list(csv.DictReader(UNION_PATH.open(encoding='utf-8-sig')))
    filled = 0
    for row in rows:
        if row['decision'] not in ('id', 'id_dot1', 'venue', 'coord'):
            continue
        if locate(row['k25_name'], row['k25_ballot']):
            continue
        donor = locate(row['k26_name'], row['k26_ballot'])
        if donor is None:
            continue
        name = normalize_name(row['k25_name'])
        ballot = row['k25_ballot']
        for key, ballot_text in ((f'{name}|{ballot}', ballot), (f'{name}|{ballot}.0', f'{ballot}.0')):
            entry = stations.get(key)
            if usable(entry):
                continue
            if entry is None:
                entry = {
                    'settlement': name,
                    'ballot': ballot_text,
                    'location': row['k25_venue'] or row['k26_venue'],
                }
                stations[key] = entry
            entry['prev_source'] = entry.get('source')
            entry['prev_lat'] = entry.get('lat')
            entry['prev_lng'] = entry.get('lng')
            entry['lat'] = donor['lat']
            entry['lng'] = donor['lng']
            entry['source'] = 'cec_k26'
            if not entry.get('location'):
                entry['location'] = row['k26_venue']
            filled += 1
    COORDS_PATH.write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(f'Backfilled {filled} coordinate keys from the paired K26 point')


def main():
    import sys
    if '--sync-keys' in sys.argv:
        sync_ballot_key_twins()
        return
    if '--backfill-missing' in sys.argv:
        backfill_missing_k25_pins()
        return
    print('Loading…')
    k25 = load_k25()
    k26 = load_k26()
    pairs, rejected, _, _ = build_pairs(k25, k26)
    proposals = correction_candidates(pairs)
    print('Scanning K16–K24 for shared coordinate keys…')
    earlier = earlier_elections(proposals)
    n = write_union(k25, k26, pairs, rejected)
    write_preview(proposals, earlier)
    print(f'Wrote {UNION_PATH} ({n} rows)')
    print(f'Wrote {PREVIEW_PATH} ({len(proposals)} rows)')
    summarize(k25, k26, pairs, rejected, proposals, earlier)
    if '--apply-geo' in sys.argv:
        apply_geo(proposals, k25, k26)
    else:
        print('Coordinates not modified (pass --apply-geo to write them)')


if __name__ == '__main__':
    main()
